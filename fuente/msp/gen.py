"""Convierte las hojas oficiales MSP (xlsx) en plantillas HTML imprimibles A4.
Uso: python3 -I gen.py <dir_src> <salida.json>
Cada celda lleva data-a="B12"; las celdas vacías llevan class "f" (rellenables)."""
import sys, json, re, html, math, colorsys
import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries
import warnings; warnings.filterwarnings("ignore")

SRC, OUT = sys.argv[1], sys.argv[2]

# hojas a incluir por formulario (None = todas las visibles)
SHEETS = {
    "007": ["SOLICITUD (1)", "INFORME (1)"],
    "016": ["ANVERSO ", "REVERSO", "FORM 016-A TESCAPURRO - B"],
    "001": ["1", "2"],
    "EPI1": ["FICHA DE NOTIFICACIÒN ", "Investigacion "],
}
BREAKS = {"EPI1": {"Investigacion ": [78]}}
# restos de datos de pacientes/establecimientos ajenos que venían en las plantillas recibidas
CLEAR = {
    "007": {"*": "A3 E3 M3 U3 A6 F6 L6 P6 T6 V6 W6 X6 Y6 Z6"},
    "017": {"017 ANVERSO": "A3 O3 U3 AJ3 BB3"},
    "EPI1": {"Investigacion ": "T6 O8 Z8 AM8 AY11"},
    "118": {"SOLICITUD": "A3 P3 V3 AT3 BP3 A6 T6 AL6 BA6 BM6 BQ6 BU6 BW6 BY6 CA6 A10 V10 AJ10 AR10 BC10 BE10 BQ10 BW10 V11 BN11 "
                         "M13 AD13 M14 AD14 BJ12 BJ13 BJ14 CA12 BP13 BP14 A16 C20 G20 M20 T20 AA20 AH20 AO20 AV20 BB20 BJ20 BU20 CA20 "
                         "H23 W23 AO23 BC23 BK23 BY23 E25 O25 AP30 AU30 AP31 AU31 AP32 AU32 AP33 AU33 AP34 AU34 AP35 AU35 AP36 AU36 "
                         "BU30 BY30 BU31 BY31 BU32 BY32 BU33 BY33 BU34 BY34 BU35 BY35 BU36 BY36 L44 AC44 AY44 BQ44"},
}
SKIP = {"028A1", "028B1", "028B2", "028C1", "028C2"}  # curvas: se generan aparte como imagen

INDEXED = ["000000","FFFFFF","FF0000","00FF00","0000FF","FFFF00","FF00FF","00FFFF",
 "000000","FFFFFF","FF0000","00FF00","0000FF","FFFF00","FF00FF","00FFFF","800000","008000","000080","808000",
 "800080","008080","C0C0C0","808080","9999FF","993366","FFFFCC","CCFFFF","660066","FF8080","0066CC","CCCCFF",
 "000080","FF00FF","FFFF00","00FFFF","800080","800000","008080","0000FF","00CCFF","CCFFFF","CCFFCC","FFFF99",
 "99CCFF","FF99CC","CC99FF","FFCC99","3366FF","33CCCC","99CC00","FFCC00","FF9900","FF6600","666699","969696",
 "003366","339966","003300","333300","993300","993366","333399","333333"]
THEME = ["FFFFFF","000000","EEECE1","1F497D","4F81BD","C0504D","9BBB59","8064A2","4BACC6","F79646"]

def tint(hexc, t):
    r, g, b = (int(hexc[i:i+2], 16)/255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l*(1+t) if t < 0 else l*(1-t)+t
    r, g, b = colorsys.hls_to_rgb(h, l, s)
    return "%02X%02X%02X" % (round(r*255), round(g*255), round(b*255))

def color(c, default=None):
    if c is None: return default
    try:
        if c.type == "rgb" and isinstance(c.rgb, str):
            v = c.rgb[-6:]
            if c.rgb.upper() in ("00000000",) and default: return default
            return "#" + v
        if c.type == "indexed":
            i = c.indexed
            if i in (64, 65): return default
            if i is not None and i < len(INDEXED): return "#" + INDEXED[i]
        if c.type == "theme" and c.theme is not None:
            base = THEME[c.theme] if c.theme < len(THEME) else "000000"
            return "#" + (tint(base, c.tint) if c.tint else base)
    except Exception:
        pass
    return default

BSTY = {"thin": "1px solid", "hair": "1px solid", "dotted": "1px dotted", "dashed": "1px dashed",
        "mediumDashed": "2px dashed", "dashDot": "1px dashed", "mediumDashDot": "2px dashed",
        "dashDotDot": "1px dotted", "slantDashDot": "2px dashed", "medium": "2px solid",
        "thick": "3px solid", "double": "3px double"}

def bcss(side):
    if side is None or side.style is None: return None
    s = BSTY.get(side.style, "1px solid")
    return s + " " + color(side.color, "#000")

def colpx(w):
    return max(0, math.floor(((256*w + math.floor(128/7))/256)*7))

STY = {}
def cls(css):
    if css not in STY: STY[css] = len(STY)
    return STY[css]
cls(""); cls("font-size:10.0pt;justify-content:flex-end;text-align:left")  # 0 y 1: celda por defecto

def is_footer(v):
    return isinstance(v, str) and "SNS-MSP" in v.upper().replace(" ", "")

def build_sheet(ws, wsf, clear, xbreaks=()):
    # extensión
    if ws.print_area:
        pa = ws.print_area if isinstance(ws.print_area, str) else ws.print_area[0]
        pa = pa.split("!")[-1].replace("$", "")
        c1, r1, c2, r2 = range_boundaries(pa)
    else:
        c1, r1, c2, r2 = 1, 1, ws.max_column, ws.max_row
    merges = {}
    covered = set()
    for m in ws.merged_cells.ranges:
        a, b, c, d = m.min_col, m.min_row, m.max_col, m.max_row
        merges[(b, a)] = (d, c)
        for rr in range(b, d+1):
            for cc in range(a, c+1):
                if (rr, cc) != (b, a): covered.add((rr, cc))
    # recorte por pie de página: hasta la última fila con SNS-MSP en columnas iniciales
    foot = [c.row for row in ws.iter_rows(min_row=r1, max_row=r2, max_col=min(c2, 6)) for c in row if is_footer(c.value)]
    if foot:
        r2 = max(foot)
        # incluir filas combinadas de la última fila
    # recorte de columnas vacías a la derecha
    used_c = c1
    for row in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2):
        for c in row:
            if c.value not in (None, "") or (c.border and (c.border.left.style or c.border.right.style or c.border.top.style or c.border.bottom.style)) or (c.fill and c.fill.patternType == "solid"):
                cc = c.column
                if (c.row, c.column) in merges: cc = merges[(c.row, c.column)][1]
                used_c = max(used_c, cc)
    c2 = min(c2, used_c)
    # anchos y altos
    dflt_w = ws.sheet_format.defaultColWidth or ((ws.sheet_format.baseColWidth or 8) + 0.71)
    cw = {}
    for k, d in ws.column_dimensions.items():
        if d.min and d.max:
            for i in range(d.min, d.max+1):
                cw[i] = (0 if d.hidden else (d.width if d.width else dflt_w))
    cols = [ci for ci in range(c1, c2+1) if cw.get(ci, dflt_w) > 0]
    colw = {ci: colpx(cw.get(ci, dflt_w)) for ci in cols}
    dflt_h = ws.sheet_format.defaultRowHeight or 15
    rows = []
    for ri in range(r1, r2+1):
        d = ws.row_dimensions.get(ri)
        if d is not None and d.hidden: continue
        h = d.height if (d is not None and d.height is not None) else dflt_h
        rows.append((ri, h*96/72))
    rowh = dict(rows)
    rowset = [r for r, _ in rows]
    colset = cols
    # páginas: cortar después de cada fila de pie
    breaks = sorted(set([r for r in foot if r < r2] + [b.id for b in ws.row_breaks.brk if r1 <= b.id < r2] + list(xbreaks)))
    pages, start = [], r1
    for b in breaks:
        pages.append((start, b)); start = b+1
    pages.append((start, r2))


    def cell_html(ri, ci, pr2):
        cell = ws.cell(ri, ci)
        er, ec = merges.get((ri, ci), (ri, ci))
        er = min(er, pr2); ec = min(ec, c2)
        rs = len([r for r in rowset if ri <= r <= er])
        cs = len([c for c in colset if ci <= c <= ec])
        if rs == 0 or cs == 0: return ""
        h = sum(rowh[r] for r in rowset if ri <= r <= er)
        w = sum(colw[c] for c in colset if ci <= c <= ec)
        # bordes del rango combinado
        def side(cellobj, name):
            return bcss(getattr(cellobj.border, name)) if cellobj.border else None
        top = next((x for x in (side(ws.cell(ri, c), "top") for c in range(ci, ec+1)) if x), None)
        bot = next((x for x in (side(ws.cell(er, c), "bottom") for c in range(ci, ec+1)) if x), None)
        lef = next((x for x in (side(ws.cell(r, ci), "left") for r in range(ri, er+1)) if x), None)
        rig = next((x for x in (side(ws.cell(r, ec), "right") for r in range(ri, er+1)) if x), None)
        css = []
        for n, v in (("top", top), ("bottom", bot), ("left", lef), ("right", rig)):
            if v: css.append("border-%s:%s" % (n, v))
        f = cell.fill
        if f is not None and f.patternType == "solid":
            bg = color(f.fgColor) or color(f.start_color)
            if bg and bg.upper() != "#FFFFFF": css.append("background:" + bg)
        fo = cell.font
        sz = (fo.sz or 10) if fo else 10
        fcss = ["font-size:%.1fpt" % sz]
        if fo and fo.b: fcss.append("font-weight:700")
        if fo and fo.i: fcss.append("font-style:italic")
        if fo and fo.u: fcss.append("text-decoration:underline")
        fc = color(fo.color, None) if fo else None
        if fc and fc.upper() not in ("#000000",): fcss.append("color:" + fc)
        al = cell.alignment
        hz = al.horizontal if al else None
        v = cell.value
        if isinstance(v, str) and v.startswith("="): v = None
        fv = wsf.cell(ri, ci).value
        if isinstance(fv, str) and fv.startswith("="): v = None
        if cell.coordinate in clear: v = None
        if hz in (None, "general"): hz = "right" if isinstance(v, (int, float)) and not isinstance(v, bool) else "left"
        if hz in ("centerContinuous", "distributed", "fill", "justify"): hz = "center" if hz != "justify" else "justify"
        vt = (al.vertical if al else None) or "bottom"
        vt = {"center": "middle", "distributed": "middle", "justify": "middle"}.get(vt, vt)
        wrap = bool(al and (al.wrap_text or al.shrink_to_fit))
        rot = al.text_rotation if al else 0
        fcss.append("justify-content:" + {"top": "flex-start", "middle": "center"}.get(vt, "flex-end"))
        c = cls(";".join(css))
        ic = cls(";".join(fcss + ["text-align:" + hz] + (["white-space:normal"] if wrap or (isinstance(v, str) and "\n" in v) else [])))
        txt = "" if v is None else (v.strftime("%d/%m/%Y") if hasattr(v, "strftime") else str(v))
        txt = txt.strip("\n")
        if txt.strip() == "": txt = ""
        fl = 0
        if wrap or "\n" in txt: fl |= 1
        if rot:
            fl |= 2 if 0 < rot <= 90 else (8 if rot == 255 else 4)
        return [ci, rs, cs, c, ic, fl, txt]

    owner = {}
    for (mr, mc), (er, ec) in merges.items():
        for rr in range(mr, er+1):
            for cc in range(mc, ec+1): owner[(rr, cc)] = (mr, mc)
    DEF_TD = cls(""); DEF_IN = cls("font-size:10.0pt;justify-content:flex-end;text-align:left")
    out_pages = []
    for (pr1, pr2) in pages:
        prow = [r for r in rowset if pr1 <= r <= pr2]
        if not prow: continue
        rows_out = []
        for ri in prow:
            cells = []
            for ci in colset:
                if (ri, ci) in covered:
                    o = owner[(ri, ci)]
                    if o[0] < pr1 and ri == pr1 and ci == o[1]:
                        x = cell_html(ri, ci, pr2)
                    else:
                        continue
                else:
                    x = cell_html(ri, ci, pr2)
                if not x: continue
                if x[1] == 1 and x[2] == 1 and x[3] == DEF_TD and x[4] == DEF_IN and x[5] == 0 and x[6] == "":
                    continue  # celda simple: la crea el JS
                x[0] = colset.index(x[0])
                while x and x[-1] in ("", 0) and len(x) > 5: x.pop()
                cells.append(x)
            rows_out.append([ri, round(rowh[ri], 1), cells])
        out_pages.append({"rows": rows_out})
    m = ws.page_margins
    return {"name": ws.title.strip(), "orient": ws.page_setup.orientation or "portrait",
            "margins": [m.top, m.right, m.bottom, m.left], "center": bool(ws.print_options.horizontalCentered),
            "cols": [get_column_letter(c) for c in colset], "cw": [colw[c] for c in colset],
            "pages": out_pages}

import glob, os
forms = {}
for path in sorted(glob.glob(os.path.join(SRC, "*.xlsx"))):
    code = os.path.basename(path)[:-5]
    if code in SKIP: continue
    wb = openpyxl.load_workbook(path, data_only=True)
    wbf = openpyxl.load_workbook(path)
    want = SHEETS.get(code)
    sheets = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible": continue
        if want and ws.title not in want: continue
        cl = CLEAR.get(code, {})
        clear = set((cl.get("*", "") + " " + cl.get(ws.title, "")).split())
        sheets.append(build_sheet(ws, wbf[ws.title], clear, BREAKS.get(code, {}).get(ws.title, ())))
    forms[code] = sheets
    print(code, [(s["name"], len(s["pages"]), sum(s["cw"])) for s in sheets])

css = "\n".join(".msp .k%d{%s}" % (k, v) for v, k in STY.items() if v)
json.dump({"css": css, "forms": forms}, open(OUT, "w"), ensure_ascii=False, separators=(",", ":"))
print("css classes", len(STY), "bytes", os.path.getsize(OUT))
